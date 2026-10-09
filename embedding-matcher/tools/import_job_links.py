#!/usr/bin/env python3
"""Import workbook job-posting and listing links into the topic-based CSV."""

from __future__ import annotations

import argparse
import csv
import re
import sys
import tempfile
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))
from runtime_paths import knowledge_base_dir

DATA_ROOT = knowledge_base_dir(PROJECT_ROOT) / "data"
WORKBOOK_PATH = DATA_ROOT / "final_master_industry_skills_dataset.xlsx"
POSTINGS_PATH = DATA_ROOT / "job_postings.csv"
ROLE_TOPIC_MAP_PATH = DATA_ROOT / "role_cluster_topic_map.csv"
OUTPUT_COLUMNS = (
    "topic",
    "posting_url",
    "posting_title",
    "employer",
    "date_retrieved",
    "notes",
    "link_type",
)


def _text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def _url_key(url: str) -> tuple[str, str, int | None, str, str]:
    parsed = urlsplit(url.strip())
    return (
        parsed.scheme.lower(),
        (parsed.hostname or "").casefold(),
        parsed.port,
        parsed.path.rstrip("/"),
        parsed.query,
    )


def _valid_http_url(value: str) -> bool:
    if not value or any(character.isspace() for character in value):
        return False
    try:
        parsed = urlsplit(value)
        if parsed.scheme.lower() not in {"http", "https"} or not parsed.hostname:
            return False
        parsed.port
        return True
    except ValueError:
        return False


def _classify_url(value: str) -> str | None:
    if not _valid_http_url(value):
        return None
    parsed = urlsplit(value)
    host = (parsed.hostname or "").lower()
    path = parsed.path.lower()
    query = parsed.query.lower()
    listing_patterns = (
        r"/search(?:/|$)",
        r"/q-[^/]*jobs?\.html?$",
        r"(?:^|[-/])jobs(?:[-/]|$)",
        r"jobs-in-",
        r"/(?:category|tag)(?:/|$)",
    )
    if any(re.search(pattern, path) for pattern in listing_patterns) or re.search(r"(?:^|&)q=", query):
        return "listing"
    if (
        re.search(r"/(?:en-us/)?job/[^/]*-\d+(?:/|$)", path)
        or (host == "jobs.recooty.com" and re.search(r"-rc\d+(?:/|$)", path))
        or "/viewjob" in path and bool(re.search(r"(?:^|&)jk=", query))
        or re.search(r"/(?:job|jobs)/\d+(?:/|$)", path)
    ):
        return "posting"
    return None


def _date_value(value: object) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raw = _text(value)
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", raw):
        try:
            return date.fromisoformat(raw).isoformat()
        except ValueError:
            return ""
    for fmt in ("%d %b %Y", "%b %d, %Y"):
        try:
            return datetime.strptime(raw.removeprefix("Posted ").strip(), fmt).date().isoformat()
        except ValueError:
            pass
    return ""


def _load_role_topics(path: Path) -> dict[str, list[str]]:
    role_topics: dict[str, list[str]] = {}
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if not {"role", "topic"}.issubset(reader.fieldnames or []):
            raise ValueError(f"Role-topic map has missing headers: {path}")
        for row in reader:
            role = _text(row.get("role")).casefold()
            topic = _text(row.get("topic")).casefold()
            if role and topic and topic not in role_topics.setdefault(role, []):
                role_topics[role].append(topic)
    return role_topics


def _row_records(workbook_path: Path, role_topics: dict[str, list[str]]):
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    rows_read = Counter()
    skipped: Counter[str] = Counter()
    skipped_details: list[dict[str, str]] = []
    records: list[dict[str, str]] = []
    sources = (
        (
            "Industry_Skills_Data",
            {
                "role": "Role_Category",
                "url": "Source_URL",
                "date": "Date_Collected",
                "source_type": "Source_Type",
            },
        ),
        (
            "Job_Postings_Log_113",
            {
                "role": "Role Cluster",
                "url": "Source Search Page (URL)",
                "date": "Posted",
                "source_type": "",
            },
        ),
    )
    for sheet_name, columns in sources:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Workbook is missing sheet {sheet_name!r}")
        sheet = workbook[sheet_name]
        headers = [_text(cell.value) for cell in sheet[1]]
        if any(columns[key] and columns[key] not in headers for key in ("role", "url", "date")):
            raise ValueError(f"Workbook sheet {sheet_name!r} has unexpected headers")
        indices = {header: index for index, header in enumerate(headers)}
        for values in sheet.iter_rows(min_row=2, values_only=True):
            if not any(value is not None and _text(value) for value in values):
                continue
            rows_read[sheet_name] += 1
            record = {header: values[index] if index < len(values) else None for header, index in indices.items()}
            role = _text(record.get(columns["role"]))
            topics = role_topics.get(role.casefold(), [])
            if not topics:
                skipped["unmapped role"] += 1
                skipped_details.append(
                    {"reason": "unmapped role", "sheet": sheet_name, "role": role, "url": ""}
                )
                continue
            url = _text(record.get(columns["url"]))
            link_type = _classify_url(url)
            if link_type is None:
                skipped["invalid URL or non-posting report/article/survey URL"] += 1
                skipped_details.append(
                    {
                        "reason": "invalid URL or non-posting report/article/survey URL",
                        "sheet": sheet_name,
                        "role": role,
                        "url": url,
                    }
                )
                continue
            parsed = urlsplit(url)
            title = f"{role} job listings ({parsed.hostname})" if link_type == "listing" else ""
            employer = ""
            if sheet_name == "Industry_Skills_Data":
                if columns["source_type"] and "job" not in _text(record.get(columns["source_type"])).casefold():
                    skipped["source is not a job posting"] += 1
                    skipped_details.append(
                        {"reason": "source is not a job posting", "sheet": sheet_name, "role": role, "url": url}
                    )
                    continue
                if link_type == "posting":
                    title = _text(record.get("Source_Name"))
            else:
                employer_value = _text(record.get("Company (if named)"))
                employer = "" if employer_value.casefold() in {"(unnamed)", "unnamed"} else employer_value
                if link_type == "posting":
                    title = _text(record.get("Posting Title / Snippet"))
            for topic in topics:
                records.append(
                    {
                        "topic": topic,
                        "posting_url": url,
                        "posting_title": title,
                        "employer": employer,
                        "date_retrieved": _date_value(record.get(columns["date"])),
                        "notes": "",
                        "link_type": link_type,
                    }
                )
    workbook.close()
    return rows_read, skipped, skipped_details, records


def import_job_links(
    workbook_path: Path = WORKBOOK_PATH,
    postings_path: Path = POSTINGS_PATH,
    role_topic_map_path: Path = ROLE_TOPIC_MAP_PATH,
) -> dict[str, object]:
    role_topics = _load_role_topics(role_topic_map_path)
    rows_read, skipped, skipped_details, candidates = _row_records(workbook_path, role_topics)
    if postings_path.exists():
        with postings_path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            existing_headers = list(reader.fieldnames or [])
            required = set(OUTPUT_COLUMNS) - {"link_type"}
            if not required.issubset(existing_headers):
                raise ValueError(f"Postings CSV has missing headers: {postings_path}")
            old_rows = list(reader)
    else:
        existing_headers = list(OUTPUT_COLUMNS)
        old_rows = []
    output_headers = existing_headers + (["link_type"] if "link_type" not in existing_headers else [])
    seen_urls = {
        (row.get("topic", "").casefold(), _url_key(_text(row.get("posting_url"))))
        for row in old_rows
        if _valid_http_url(_text(row.get("posting_url")))
    }
    added: list[dict[str, str]] = []
    for candidate in candidates:
        key = (candidate["topic"].casefold(), _url_key(candidate["posting_url"]))
        if key in seen_urls:
            skipped["duplicate URL"] += 1
            continue
        seen_urls.add(key)
        added.append(candidate)
    if added or output_headers != existing_headers:
        postings_path.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(
            "w",
            encoding="utf-8-sig",
            newline="",
            dir=postings_path.parent,
            delete=False,
        ) as handle:
            temp_path = Path(handle.name)
            writer = csv.DictWriter(handle, fieldnames=output_headers, extrasaction="ignore")
            writer.writeheader()
            for row in old_rows:
                if not row.get("link_type"):
                    row["link_type"] = "posting"
                writer.writerow(row)
            for row in added:
                writer.writerow(row)
        temp_path.replace(postings_path)

    added_topic_counts: Counter[tuple[str, str]] = Counter(
        (row["topic"], row["link_type"]) for row in added
    )
    total_topic_counts: Counter[tuple[str, str]] = Counter()
    for row in old_rows + added:
        topic = _text(row.get("topic")).casefold()
        link_type = _text(row.get("link_type")).casefold() or "posting"
        if topic and link_type in {"posting", "listing"}:
            total_topic_counts[(topic, link_type)] += 1
    all_topics = sorted({topic for topics in role_topics.values() for topic in topics})
    return {
        "rows_read": dict(rows_read),
        "kept_posting": sum(row["link_type"] == "posting" for row in added),
        "kept_listing": sum(row["link_type"] == "listing" for row in added),
        "skipped": dict(skipped),
        "skipped_details": skipped_details,
        "per_topic": {
            (topic, link_type): total_topic_counts[(topic, link_type)]
            for topic in all_topics
            for link_type in ("posting", "listing")
        },
        "added_per_topic": dict(added_topic_counts),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=WORKBOOK_PATH)
    parser.add_argument("--postings", type=Path, default=POSTINGS_PATH)
    parser.add_argument("--role-topic-map", type=Path, default=ROLE_TOPIC_MAP_PATH)
    args = parser.parse_args()
    try:
        summary = import_job_links(args.workbook, args.postings, args.role_topic_map)
    except Exception as error:
        print(f"Job-link import failed: {error}", file=sys.stderr)
        return 1
    print("Rows read:")
    for sheet, count in summary["rows_read"].items():
        print(f"- {sheet}: {count}")
    print(f"Kept as postings: {summary['kept_posting']}")
    print(f"Kept as listings: {summary['kept_listing']}")
    print("Skipped:")
    for reason, count in summary["skipped"].items():
        print(f"- {reason}: {count}")
    print("Skipped source rows:")
    for detail in summary["skipped_details"]:
        if detail["reason"] == "unmapped role":
            print(f"- {detail['sheet']}: unmapped role {detail['role']!r}")
        elif detail["url"]:
            print(f"- {detail['sheet']} ({detail['role']}): {detail['url']} [{detail['reason']}]")
    print("Current total links by topic:")
    for topic, link_type in sorted(summary["per_topic"]):
        print(f"- {topic}: {summary['per_topic'][(topic, link_type)]} {link_type}(s)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
